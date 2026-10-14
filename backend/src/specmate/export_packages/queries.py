"""Coherent export packages: one manifest pinning every included output to exact records.

Existing generators keep their own semantics and authority. This module resolves one
exact root closure, independently counts contextual occurrences for the stated scope
quantity, and checks that every retained output cites the same intended revisions
before packaging its exact bytes. Nothing here runs a native tool or writes state.
"""

import csv
import hashlib
import html
import io
import json
import zipfile
from collections import Counter, defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.writer.excel import ExcelWriter

from specmate.assemblies.queries import get_snapshot, occurrence_result
from specmate.bom.queries import inspect_revision as inspect_bom
from specmate.cable_test_intent.queries import portable_package as test_package
from specmate.cable_test_intent.queries import result as test_result
from specmate.cables.queries import inspect_revision as inspect_cable
from specmate.catalog.models import Representation, SourceIdentifier
from specmate.catalog_matching.queries import current_actor
from specmate.design.models import Revision
from specmate.design.operations import canonical_bytes
from specmate.harnesses.queries import inspect_revision as inspect_harness
from specmate.operations import Conflict, ValidationFailed
from specmate.pcb.models import PcbInterpretation
from specmate.pcb_fabrication.operations import scoped_build, validate_outputs
from specmate.pcb_fabrication.pack import REPORT_PATH
from specmate.preparation_documents.operations import authorized_set
from specmate.preparation_documents.rendering import MAX_FILE_BYTES
from specmate.queries import get_revision
from specmate.shared_scene.drawing_operations import drawing_record
from specmate.sources import blobstore

FORMAT = "specmate.coherent-export-package"
SCHEMA_VERSION = "0.1"
GENERATOR = "specmate.export-packages/1"
XLSX_PATH = "spreadsheet/package.xlsx"
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_BYTES = 16 * 1024 * 1024
MAX_FILES = 4000
MAX_REVISIONS = 2000
EPOCH = (1980, 1, 1, 0, 0, 0)
WORKBOOK_TIME = datetime(1980, 1, 1)

# Output semantics this package does not generate. They are stated, never faked.
UNSUPPORTED = [
    {
        "id": "cable.formboard_print",
        "section": "cable",
        "reason": "Harness formboard/print views are rendered in the browser and are not "
        "retained server outputs; they are not embedded.",
    },
    {
        "id": "pcb.assembly_outputs",
        "section": "pcb",
        "reason": "Pick-and-place, population BOM and assembly drawings are not produced by "
        "the registered fabrication pack.",
    },
    {
        "id": "mechanical.native_cad_export",
        "section": "mechanical",
        "reason": "No STEP/native CAD or toleranced manufacturing drawing is generated. "
        "Preserved native mechanical sources are cited by exact pin, not re-exported.",
    },
    {
        "id": "mechanical.scaled_drawing",
        "section": "mechanical",
        "reason": "Retained scene drawings are internal inspection views; scale and "
        "calibration are not established.",
    },
    {
        "id": "firmware.build_outputs",
        "section": "system",
        "reason": "Firmware revisions are pinned; compiled firmware outputs are not "
        "packaged by this generator.",
    },
    {
        "id": "test.results",
        "section": "test",
        "reason": "Packages carry authored test intent only; no physical test was performed.",
    },
    {
        "id": "distribution",
        "section": "package",
        "reason": "Only authorized internal review is supported. The rights review lists "
        "declared catalog rights; it is not a distribution grant or legal clearance.",
    },
]

SECTIONS = ("system", "cable", "pcb", "mechanical", "spreadsheet", "test", "rights")


def digest(data):
    return hashlib.sha256(data).hexdigest()


def cell(value):
    """CSV text never becomes a spreadsheet formula; unknown stays an empty cell."""
    if value is None:
        return ""
    text = str(value)
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else text


def table(header, rows):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, lineterminator="\r\n")
    writer.writerow(header)
    for row in rows:
        writer.writerow([cell(value) for value in row])
    return stream.getvalue().encode()


def plain(value):
    return value if value is None or isinstance(value, bool | int | float) else str(value)


def workbook(sheets):
    """Deterministic XLSX of `(name, header, rows)`; same input gives identical bytes.

    Text is always stored as a string cell, so it can never become a formula. Unknown
    values are empty cells. The core-properties timestamps are fixed and the zip members
    are rewritten with a fixed timestamp in sorted order.
    """
    book = Workbook()
    book.remove(book.active)
    book.properties.creator = GENERATOR
    book.properties.lastModifiedBy = None
    book.properties.created = WORKBOOK_TIME
    book.properties.modified = WORKBOOK_TIME
    for name, header, rows in sheets:
        sheet = book.create_sheet(name[:31])
        sheet.append(list(header))
        for row in rows:
            sheet.append(
                [
                    None if v is None else v if isinstance(v, bool | int | float) else str(v)
                    for v in row
                ]
            )
        for line in sheet.iter_rows():
            for item in line:
                if isinstance(item.value, str):
                    item.data_type = "s"
        sheet.freeze_panes = "A2"
    raw = io.BytesIO()
    # openpyxl's `save_workbook` stamps `modified` with the current time; its writer does not.
    with zipfile.ZipFile(raw, "w", zipfile.ZIP_DEFLATED) as archive:
        ExcelWriter(book, archive).save()
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as source:
        with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for name in sorted(source.namelist()):
                info = zipfile.ZipInfo(name, date_time=EPOCH)
                info.create_system = 3
                info.external_attr = 0o100644 << 16
                info.compress_type = zipfile.ZIP_DEFLATED
                target.writestr(info, source.read(name))
    return out.getvalue()


def pin(revision):
    return {
        "revision_id": str(revision.id),
        "definition_id": str(revision.definition_id),
        "kind": revision.definition.kind,
        "name": revision.definition.name,
        "number": revision.number,
        "snapshot_id": str(revision.snapshot_id),
        "sha256": revision.snapshot.sha256,
    }


def scope_quantity(count, quantity):
    """Per-root occurrence count times stated root quantity; unknown is never zero."""
    if count is None:
        return {"status": "not_applicable", "amount": None}
    if quantity is None:
        return {"status": "unknown", "amount": None, "per_root": count}
    return {"status": "known", "amount": count * quantity, "per_root": count}


class Scope:
    """The exact root closure. Every pin is read through the shared authorized reads."""

    def __init__(self, user, p):
        self.user = user
        self.quantity = p.quantity
        self.root = get_revision(user, p.root_revision_id)
        self.closure = None
        self.revisions = {self.root.id: self.root}
        if self.root.definition.kind == "assembly":
            _, self.closure = get_snapshot(user, self.root.snapshot_id)
            self.revisions.update(self.closure.revisions)
        else:
            pending = list(self.root.snapshot.dependencies.values_list("revision_id", flat=True))
            while pending:
                rid = pending.pop()
                if rid in self.revisions:
                    continue
                if len(self.revisions) >= MAX_REVISIONS:
                    raise ValidationFailed("Export package dependency limit exceeded.")
                revision = get_revision(user, rid)
                self.revisions[rid] = revision
                pending.extend(revision.snapshot.dependencies.values_list("revision_id", flat=True))
        if len(self.revisions) > MAX_REVISIONS:
            raise ValidationFailed("Export package dependency limit exceeded.")
        # Independent structural count: the root once, then each contextual occurrence.
        self.counts = Counter({self.root.id: 1})
        self.paths = defaultdict(list)
        self.paths[self.root.id].append([])
        if self.closure:
            for path, node in self.closure.nodes.items():
                self.counts[node.revision.id] += 1
                self.paths[node.revision.id].append(list(path))
        self.definitions = defaultdict(set)
        for revision in self.revisions.values():
            self.definitions[revision.definition_id].add(revision.id)
        self.snapshots = {r.snapshot_id: r for r in self.revisions.values()}

    def count(self, rid):
        return self.counts.get(rid)

    def scope_quantity(self, rid):
        return scope_quantity(self.count(rid), self.quantity)

    def of_kind(self, *kinds):
        return sorted(
            (r for r in self.revisions.values() if r.definition.kind in kinds),
            key=lambda r: str(r.id),
        )

    def cite(self, revision_id, definition_id):
        """`pinned` same revision; `conflict` another revision of a pinned definition."""
        pinned = self.definitions.get(definition_id)
        if pinned is None:
            status = "supplementary"
        elif revision_id in pinned:
            status = "pinned"
        else:
            status = "conflict"
        return {
            "revision_id": str(revision_id),
            "definition_id": str(definition_id),
            "status": status,
            "package_revision_ids": sorted(str(r) for r in pinned or ()),
        }


class Package:
    def __init__(self):
        self.files = {}
        self.media = {}
        self.sections = {}
        self.outputs = []
        self.unresolved = []

    def add(self, path, data, media_type, section, origin):
        if path in self.files:
            raise ValidationFailed("Duplicate export package path.")
        if len(self.files) >= MAX_FILES or sum(map(len, self.files.values())) + len(data) > (
            MAX_BYTES
        ):
            raise ValidationFailed("Export package exceeds the supported size limit.")
        self.files[path] = data
        self.media[path] = (media_type, section, origin)


def output_record(section, kind, ident, anchor, cited, findings, quantity=None, **extra):
    statuses = {f["status"] for f in findings}
    coherence = "incoherent" if "incoherent" in statuses else "unknown" if statuses else "coherent"
    return {
        "id": f"{kind}:{ident}",
        "section": section,
        "kind": kind,
        "anchor": anchor,
        "cited_revisions": cited,
        "scope_quantity": quantity,
        "coherence": coherence,
        "findings": findings,
        **extra,
    }


def conflicts(cited):
    return [
        {
            "status": "incoherent",
            "reason": "cites_different_revision",
            "revision_id": row["revision_id"],
            "definition_id": row["definition_id"],
            "package_revision_ids": row["package_revision_ids"],
        }
        for row in cited
        if row["status"] == "conflict"
    ]


def system_section(scope, package):
    root = scope.root
    occurrences = []
    relationships = []
    if scope.closure:
        occurrences = [
            {
                **occurrence_result(root.snapshot, node, scope.closure.assignments.get(path)),
                "per_root_count": 1,
            }
            for path, node in scope.closure.nodes.items()
        ]
        relationships = [
            {"context_path": list(prefix), **rel.model_dump(mode="json")}
            for prefix, rel in scope.closure.relationships
        ]
    structure = {
        "root": pin(root),
        "quantity": scope.quantity,
        "occurrences": occurrences,
        "relationships": relationships,
        "firmware_revisions": [pin(r) for r in scope.of_kind("firmware")],
        "connectivity": "declared_only",
        "physical_qualification": "not_evaluated",
    }
    package.add(
        "system/structure.json",
        canonical_bytes(structure),
        "application/json",
        "system",
        "generated",
    )
    for rid in sorted(scope.revisions, key=str):
        revision = scope.revisions[rid]
        raw = bytes(revision.snapshot.canonical)
        if digest(raw) != revision.snapshot.sha256:
            raise ValidationFailed("Pinned design bytes failed integrity verification.")
        package.add(f"designs/{rid}.json", raw, "application/json", "system", "canonical")
    package.sections["system"] = {
        "status": "included",
        "occurrences": len(occurrences),
        "pinned_revisions": len(scope.revisions),
    }


OCCURRENCE_HEADER = [
    "path",
    "display_ref",
    "kind",
    "name",
    "definition_id",
    "revision_number",
    "revision_id",
    "snapshot_sha256",
    "per_root_count",
    "scope_quantity",
    "scope_quantity_status",
]
QUANTITY_HEADER = [
    "revision_id",
    "definition_id",
    "kind",
    "name",
    "revision_number",
    "snapshot_sha256",
    "role",
    "per_root_occurrences",
    "scope_quantity",
    "scope_quantity_status",
]
BOM_HEADER = [
    "path",
    "name",
    "revision_id",
    "reason",
    "need_id",
    "quantity",
    "quantity_unit",
    "quantity_status",
    "required_quantity",
    "required_unit",
    "required_status",
]


def spreadsheet_section(scope, package, p):
    sheets = []

    def add_table(path, name, header, rows, origin="generated"):
        package.add(path, table(header, rows), "text/csv", "spreadsheet", origin)
        sheets.append((name, header, rows))

    rows = []
    if scope.closure:
        for path, node in scope.closure.nodes.items():
            revision = node.revision
            quantity = scope_quantity(1, scope.quantity)
            rows.append(
                [
                    "/".join(path),
                    node.occurrence.display_ref,
                    revision.definition.kind,
                    revision.definition.name,
                    revision.definition_id,
                    revision.number,
                    revision.id,
                    revision.snapshot.sha256,
                    1,
                    quantity["amount"],
                    quantity["status"],
                ]
            )
    add_table("spreadsheet/occurrences.csv", "occurrences", OCCURRENCE_HEADER, rows)
    rows = []
    for rid in sorted(scope.revisions, key=str):
        revision = scope.revisions[rid]
        count = scope.count(rid)
        quantity = scope.scope_quantity(rid)
        role = "root" if rid == scope.root.id else "occurrence" if count else "dependency"
        rows.append(
            [
                rid,
                revision.definition_id,
                revision.definition.kind,
                revision.definition.name,
                revision.number,
                revision.snapshot.sha256,
                role,
                count,
                quantity["amount"],
                quantity["status"],
            ]
        )
    add_table("spreadsheet/revision-quantities.csv", "revision-quantities", QUANTITY_HEADER, rows)
    section = {"status": "included", "bom": "not_included", "workbook": XLSX_PATH}
    if p.bom_plan_revision_id:
        bom = inspect_bom(scope.user, p.bom_plan_revision_id, p.bom_purpose)
        plan = get_revision(scope.user, p.bom_plan_revision_id)
        cited = [scope.cite(plan.id, plan.definition_id)]
        cited += [
            scope.cite(r.id, r.definition_id)
            for r in Revision.objects.select_related("definition").filter(
                id__in=plan.snapshot.dependencies.values_list("revision_id", flat=True)
            )
        ]
        findings = conflicts(cited)
        if bom["bom"]["assembly_revision_id"] != str(scope.root.id):
            findings.append({"status": "incoherent", "reason": "bom_root_differs"})
        declared = bom["bom"]["quantity"]
        amount = declared.get("amount") if declared.get("status") == "known" else None
        if amount != scope.quantity:
            findings.append(
                {
                    "status": "incoherent",
                    "reason": "bom_quantity_differs",
                    "bom_quantity": amount,
                    "package_quantity": scope.quantity,
                }
            )
        path = f"spreadsheet/bom-{p.bom_purpose}.csv"
        add_table(
            path,
            f"bom-{p.bom_purpose}"[:31],
            BOM_HEADER,
            [
                [
                    "/".join(row["address"]["path"]),
                    row["name"],
                    row["revision_id"],
                    row["reason"],
                    row["need_id"],
                    row["quantity"]["amount"],
                    row["quantity"]["unit"],
                    row["quantity"]["status"],
                    row["required_quantity"]["amount"],
                    row["required_quantity"]["unit"],
                    row["required_quantity"]["status"],
                ]
                for row in bom["rows"]
            ],
            "bom_projection",
        )
        package.add(
            f"spreadsheet/bom-{p.bom_purpose}.json",
            canonical_bytes(bom),
            "application/json",
            "spreadsheet",
            "bom_projection",
        )
        for item in bom["unresolved"]:
            package.unresolved.append({"output": f"bom_plan:{plan.id}", **item})
        package.outputs.append(
            output_record(
                "spreadsheet",
                "bom_plan",
                plan.id,
                {"root_revision_id": bom["bom"]["assembly_revision_id"]},
                cited,
                findings,
                {"status": "known" if amount else "unknown", "amount": amount},
                purpose=p.bom_purpose,
                status=bom["status"],
                files=[path, f"spreadsheet/bom-{p.bom_purpose}.json"],
            )
        )
        section["bom"] = "included"
    package.add(XLSX_PATH, workbook(sheets), XLSX_MEDIA, "spreadsheet", "workbook")
    package.sections["spreadsheet"] = section


def cable_section(scope, package, p):
    cables = scope.of_kind("cable")
    harnesses = scope.of_kind("harness")
    for revision in cables:
        inspected = inspect_cable(scope.user, revision.id)
        base = f"cable/{revision.id}"
        package.add(
            f"{base}/cable.json", canonical_bytes(inspected), "application/json", "cable", "cable"
        )
        rows = []
        for conductor in inspected["conductors"]:
            for end in conductor["ends"]:
                assignment = end["assignment"]
                rows.append(
                    [
                        conductor["element_path"],
                        conductor["kind"],
                        conductor["label"],
                        end["endpoint_id"],
                        assignment.get("contact_id") if assignment else None,
                        "assigned" if assignment else "unassigned",
                    ]
                )
        package.add(
            f"{base}/conductors.csv",
            table(["element_path", "kind", "label", "endpoint_id", "contact_id", "status"], rows),
            "text/csv",
            "cable",
            "cable",
        )
    for revision in harnesses:
        package.add(
            f"cable/{revision.id}/harness.json",
            canonical_bytes(inspect_harness(scope.user, revision.id)),
            "application/json",
            "cable",
            "harness",
        )
    for set_id in p.preparation_document_set_ids:
        record = authorized_set(scope.user, set_id)
        inputs = record.manifest["inputs"]
        cited = [
            scope.cite(row["revision_id"], row["definition_id"]) for row in inputs["revision_pins"]
        ]
        findings = conflicts(cited)
        root_id = record.root_revision_id
        expected = None
        if root_id not in scope.revisions:
            findings.append({"status": "incoherent", "reason": "document_root_not_in_package"})
        else:
            expected = scope.scope_quantity(root_id)
            declared = record.configuration.get("quantity")
            if expected["status"] == "known" and declared != expected["amount"]:
                findings.append(
                    {
                        "status": "incoherent",
                        "reason": "document_quantity_differs",
                        "document_quantity": declared,
                        "expected_quantity": expected["amount"],
                    }
                )
            elif expected["status"] != "known":
                findings.append(
                    {
                        "status": "unknown",
                        "reason": "scope_quantity_unknown",
                        "document_quantity": declared,
                    }
                )
        bom_id = record.configuration.get("bom_plan_revision_id")
        if bom_id and root_id == scope.root.id and p.bom_plan_revision_id is not None:
            if bom_id != str(p.bom_plan_revision_id):
                findings.append({"status": "incoherent", "reason": "document_bom_differs"})
        paths = []
        for output in sorted(record.outputs.all(), key=lambda o: o.filename):
            with blobstore.open_blob(output.sha256) as stream:
                data = stream.read(MAX_FILE_BYTES + 1)
            if len(data) != output.size or digest(data) != output.sha256:
                raise ValidationFailed("Retained document integrity mismatch.")
            path = f"cable/preparation/{record.id}/{output.filename}"
            package.add(path, data, output.media_type, "cable", "retained_document_set")
            paths.append(path)
        package.outputs.append(
            output_record(
                "cable",
                "preparation_document_set",
                record.id,
                {"root_revision_id": str(root_id), "inputs_sha256": record.sha256},
                cited,
                findings,
                expected,
                generator=record.generator,
                files=paths,
            )
        )
    package.sections["cable"] = {
        "status": "included" if cables or harnesses or p.preparation_document_set_ids else "empty",
        "cables": len(cables),
        "harnesses": len(harnesses),
        "document_sets": len(p.preparation_document_set_ids),
    }


def test_section(scope, package, p):
    for rid in p.test_plan_revision_ids:
        inspected = test_result(scope.user, rid)
        data, sha256 = test_package(scope.user, rid)
        plan = inspected["test_plan"]
        revision = get_revision(scope.user, rid)
        cited = [scope.cite(revision.id, revision.definition_id)] + [
            scope.cite(r.id, r.definition_id)
            for r in Revision.objects.select_related("definition").filter(
                id__in=revision.snapshot.dependencies.values_list("revision_id", flat=True)
            )
        ]
        findings = conflicts(cited)
        root_id = plan["root_revision_id"]
        anchors = []
        quantity = None
        roots = [r for r in scope.revisions.values() if str(r.id) == root_id]
        if not roots:
            findings.append({"status": "incoherent", "reason": "test_root_not_in_package"})
        else:
            # A plan addressed below its own root applies once per occurrence of that root.
            root = roots[0]
            for prefix in scope.paths.get(root.id, []):
                anchors.append([*prefix, *plan["occurrence_path"]])
            quantity = scope.scope_quantity(root.id)
            if quantity["status"] == "not_applicable":
                findings.append({"status": "unknown", "reason": "test_root_not_an_occurrence"})
        path = f"test/{rid}/cable-test-package.json"
        if digest(data) != sha256:
            raise ValidationFailed("Cable test package integrity mismatch.")
        package.add(path, data, "application/json", "test", "cable_test_package")
        package.outputs.append(
            output_record(
                "test",
                "cable_test_plan",
                rid,
                {
                    "root_revision_id": root_id,
                    "cable_revision_id": plan["cable_revision_id"],
                    "occurrence_paths": anchors,
                },
                cited,
                findings,
                quantity,
                physical_test="not_performed",
                files=[path],
            )
        )
    package.sections["test"] = {
        "status": "included" if p.test_plan_revision_ids else "empty",
        "test_plans": len(p.test_plan_revision_ids),
    }


def population_sources(scope):
    """Exact board source revision/path pinned by each population revision in scope."""
    rows = []
    for revision in scope.of_kind("pcb_population"):
        content = json.loads(bytes(revision.snapshot.canonical))["content"]
        interpretation_id = content.get("pcb_population", {}).get("interpretation_id")
        interpretation = (
            PcbInterpretation.objects.select_related("attempt__intake")
            .filter(id=interpretation_id)
            .first()
        )
        intake = interpretation.attempt.intake if interpretation else None
        rows.append(
            {
                "revision": revision,
                "interpretation_id": interpretation_id,
                "intake_id": str(intake.id) if intake else None,
                "source_revision_id": str(intake.source_revision_id) if intake else None,
                "board_path": intake.board_path if intake else None,
            }
        )
    return rows


def pcb_section(ctx, scope, package, p):
    populations = population_sources(scope)
    built = set()
    for ref in p.pcb_fabrication_builds:
        intake, file, build = scoped_build(ctx, ref.intake_id, ref.build_id)
        report, outputs = validate_outputs(file, build)
        matches = [
            row
            for row in populations
            if row["source_revision_id"] == str(file.revision_id)
            and row["board_path"] == intake.board_path
        ]
        cited = [scope.cite(row["revision"].id, row["revision"].definition_id) for row in matches]
        findings = conflicts(cited)
        quantity = None
        if not matches:
            findings.append({"status": "incoherent", "reason": "board_source_not_pinned"})
        else:
            counts = [scope.count(row["revision"].id) for row in matches]
            per_root = sum(c or 0 for c in counts)
            quantity = scope_quantity(per_root or None, scope.quantity)
            if not per_root:
                findings.append({"status": "unknown", "reason": "board_not_an_occurrence"})
        base = f"pcb/{build.id}"
        paths = []
        for path, output in sorted(outputs.items()):
            with blobstore.open_blob(output.sha256) as stream:
                data = stream.read(MAX_BYTES + 1)
            if len(data) != output.size or digest(data) != output.sha256:
                raise ValidationFailed("Stored fabrication bytes differ from registered evidence.")
            target = f"{base}/{path}"
            package.add(target, data, output.media_type, "pcb", "registered_build_output")
            paths.append(target)
            built.update(row["revision"].id for row in matches)
        package.outputs.append(
            output_record(
                "pcb",
                "pcb_fabrication_build",
                build.id,
                {
                    "intake_id": str(intake.id),
                    "source_revision_id": str(file.revision_id),
                    "board_path": intake.board_path,
                    "source_sha256": file.blob.sha256,
                    "recipe_id": build.recipe_id,
                    "report_sha256": outputs[REPORT_PATH].sha256,
                },
                cited,
                findings,
                quantity,
                drc_counts=report.drc_counts.model_dump(),
                manufacturing_approval=False,
                files=paths,
            )
        )
    for row in populations:
        if row["revision"].id not in built:
            package.unresolved.append(
                {
                    "output": "pcb",
                    "revision_id": str(row["revision"].id),
                    "reason": "no_registered_fabrication_output_included",
                }
            )
    package.add(
        "pcb/populations.json",
        canonical_bytes(
            [
                {
                    **pin(row["revision"]),
                    "interpretation_id": row["interpretation_id"],
                    "source_revision_id": row["source_revision_id"],
                    "board_path": row["board_path"],
                    "scope_quantity": scope.scope_quantity(row["revision"].id),
                }
                for row in populations
            ]
        ),
        "application/json",
        "pcb",
        "generated",
    )
    package.sections["pcb"] = {
        "status": "included" if populations or p.pcb_fabrication_builds else "empty",
        "populations": len(populations),
        "fabrication_builds": len(p.pcb_fabrication_builds),
    }


def mechanical_section(scope, package, p):
    sources = scope.of_kind("mechanical_source")
    for drawing_id in p.scene_drawing_ids:
        drawing = drawing_record(scope.user, drawing_id)
        revision = scope.snapshots.get(drawing.snapshot_id)
        cited = [scope.cite(revision.id, revision.definition_id)] if revision else []
        findings = []
        if revision is None:
            findings.append({"status": "incoherent", "reason": "drawing_scene_not_in_package"})
        base = f"mechanical/{drawing.id}"
        package.add(
            f"{base}/drawing.svg", bytes(drawing.svg), "image/svg+xml", "mechanical", "drawing"
        )
        package.add(
            f"{base}/manifest.json",
            bytes(drawing.manifest),
            "application/json",
            "mechanical",
            "drawing",
        )
        package.outputs.append(
            output_record(
                "mechanical",
                "scene_drawing",
                drawing.id,
                {
                    "snapshot_id": str(drawing.snapshot_id),
                    "renderer_version": drawing.renderer_version,
                    "purpose": drawing.purpose,
                },
                cited,
                findings,
                scope.scope_quantity(revision.id) if revision else None,
                scale="not_established",
                files=[f"{base}/drawing.svg", f"{base}/manifest.json"],
            )
        )
    package.add(
        "mechanical/sources.json",
        canonical_bytes(
            [{**pin(r), "scope_quantity": scope.scope_quantity(r.id)} for r in sources]
        ),
        "application/json",
        "mechanical",
        "generated",
    )
    package.sections["mechanical"] = {
        "status": "included" if sources or p.scene_drawing_ids else "empty",
        "mechanical_sources": len(sources),
        "native_bytes": "not_included",
        "drawings": len(p.scene_drawing_ids),
    }


RIGHTS_HEADER = [
    "subject",
    "revision_id",
    "definition_name",
    "record",
    "label",
    "rights_status",
    "declared_rights",
    "source_path",
    "sha256",
]


def rights_section(scope, package):
    """Declared catalog rights of every representation and source identifier in the closure.

    Rights status is read as recorded (cleared, restricted, unknown). Nothing is inferred
    and a revision without any declaration is listed as undeclared, never as cleared.
    """
    ids = list(scope.revisions)
    representations = (
        Representation.objects.select_related(
            "revision__definition", "source_file__blob", "source_file__revision__bundle"
        )
        .filter(revision_id__in=ids)
        .order_by("id")
    )
    items = []
    for item in representations:
        # Same workspace boundary as the catalog read: a file of another workspace is hidden.
        if item.source_file.revision.bundle.workspace_id != item.revision.definition.workspace_id:
            continue
        items.append(
            {
                "subject": f"representation:{item.id}",
                "revision_id": str(item.revision_id),
                "definition_name": item.revision.definition.name,
                "record": item.kind,
                "label": item.label,
                "rights_status": item.rights_status,
                "declared_rights": item.source_file.rights,
                "source_path": item.source_file.path,
                "sha256": item.source_file.blob.sha256,
            }
        )
    identifiers = (
        SourceIdentifier.objects.select_related("revision__definition")
        .filter(revision_id__in=ids)
        .order_by("id")
    )
    for item in identifiers:
        items.append(
            {
                "subject": f"source_identifier:{item.id}",
                "revision_id": str(item.revision_id),
                "definition_name": item.revision.definition.name,
                "record": item.scheme,
                "label": item.label or item.identifier,
                "rights_status": item.rights_status,
                "declared_rights": item.rights,
                "source_path": None,
                "sha256": None,
            }
        )
    items.sort(key=lambda row: row["subject"])
    counts = Counter(row["rights_status"] for row in items)
    declared = {row["revision_id"] for row in items}
    undeclared = sorted(str(rid) for rid in ids if str(rid) not in declared)
    review = {
        "items": items,
        "counts": {name: counts.get(name, 0) for name in ("cleared", "restricted", "unknown")},
        "undeclared_revision_ids": undeclared,
        "acknowledgement_required": any(row["rights_status"] != "cleared" for row in items),
        "distribution": "not_authorized",
    }
    package.add(
        "rights/review.json", canonical_bytes(review), "application/json", "rights", "generated"
    )
    package.add(
        "rights/review.csv",
        table(RIGHTS_HEADER, [[row[name] for name in RIGHTS_HEADER] for row in items]),
        "text/csv",
        "rights",
        "generated",
    )
    package.sections["rights"] = {
        "status": "included",
        **review["counts"],
        "undeclared_revisions": len(undeclared),
    }
    return review


def review_document(manifest):
    # Static escaped HTML; no scripts, remote resources or executable templates.
    def block(title, value):
        return f"<h2>{html.escape(title)}</h2><pre>{html.escape(json.dumps(value, indent=2))}</pre>"

    return (
        "<!doctype html><html lang='en'><meta charset='utf-8'>"
        "<meta http-equiv='Content-Security-Policy' content=\"default-src 'none'\">"
        "<title>Coherent export package</title><h1>Coherent export package</h1>"
        f"<p>Coherence: {html.escape(manifest['coherence'])}. "
        f"Completeness: {html.escape(manifest['completeness'])}. "
        f"Root quantity: {html.escape(str(manifest['scope']['quantity']))}.</p>"
        "<p>Declared engineering data. Not a release, manufacturing approval or physical "
        "qualification. Unknown quantities stay unknown.</p>"
        + block("Exact root and scope", manifest["scope"])
        + block("Sections", manifest["sections"])
        + block("Included outputs and coherence", manifest["outputs"])
        + block("Rights and distribution review", manifest["rights"])
        + block("Unresolved items", manifest["unresolved"])
        + block("Unsupported output semantics", manifest["unsupported"])
        + block("Pinned revisions", manifest["pins"])
        + "</html>"
    ).encode()


def build_package(ctx, p):
    user = current_actor(ctx.user)
    scope = Scope(user, p)
    package = Package()
    system_section(scope, package)
    spreadsheet_section(scope, package, p)
    cable_section(scope, package, p)
    pcb_section(ctx, scope, package, p)
    mechanical_section(scope, package, p)
    test_section(scope, package, p)
    rights = rights_section(scope, package)
    statuses = {o["coherence"] for o in package.outputs}
    # Coherence: do outputs agree on revisions and quantities? Completeness is separate:
    # agreeing outputs can still carry unresolved items or an unknown scope quantity.
    coherence = (
        "incoherent"
        if "incoherent" in statuses
        else "unknown"
        if "unknown" in statuses or p.quantity is None
        else "coherent"
    )
    completeness = "limited" if package.unresolved or coherence != "coherent" else "complete"
    manifest = {
        "format": FORMAT,
        "schema_version": SCHEMA_VERSION,
        "generator": GENERATOR,
        "purpose": p.purpose,
        "request": p.model_dump(mode="json", exclude={"reviewed_sha256", "rights_acknowledged"}),
        "scope": {
            "root": pin(scope.root),
            "quantity": p.quantity,
            "quantity_status": "known" if p.quantity is not None else "unknown",
        },
        "pins": [
            {
                **pin(scope.revisions[rid]),
                "per_root_occurrences": scope.count(rid),
                "scope_quantity": scope.scope_quantity(rid),
                "design_path": f"designs/{rid}.json",
            }
            for rid in sorted(scope.revisions, key=str)
        ],
        "sections": {name: package.sections.get(name, {"status": "empty"}) for name in SECTIONS},
        "outputs": package.outputs,
        "coherence": coherence,
        "completeness": completeness,
        "unresolved": package.unresolved,
        "unsupported": UNSUPPORTED,
        "rights": rights,
        "distribution": "not_authorized",
        "manufacturing_approval": False,
        "physical_qualification": "not_evaluated",
    }
    package.add("review.html", review_document(manifest), "text/html", "package", "generated")
    package.add(
        "README.txt",
        b"Specmate coherent export package (format specmate.coherent-export-package 0.1).\n"
        b"manifest.json pins every included output to exact revisions and lists each file's\n"
        b"SHA-256 and byte count. Verify files against it, and verify manifest.json against\n"
        b"the SHA-256 retained from the authorized review. Files are inert data; nothing is\n"
        b"a release, manufacturing approval, physical test or qualification.\n",
        "text/plain",
        "package",
        "generated",
    )
    manifest["files"] = [
        {
            "path": path,
            "sha256": digest(data),
            "bytes": len(data),
            "media_type": package.media[path][0],
            "section": package.media[path][1],
            "origin": package.media[path][2],
        }
        for path, data in sorted(package.files.items())
    ]
    data = canonical_bytes(manifest)
    if sum(map(len, package.files.values())) + len(data) > MAX_BYTES:
        raise ValidationFailed("Export package exceeds the supported size limit.")
    package.files["manifest.json"] = data
    return manifest, package.files, digest(data)


def archive(ctx, p):
    manifest, files, sha256 = build_package(ctx, p)
    if sha256 != p.reviewed_sha256:
        raise Conflict("Export package changed since review; review the exact request again.")
    if manifest["coherence"] == "incoherent":
        raise ValidationFailed(
            "Included outputs cite different revisions or quantities; no package was produced."
        )
    if manifest["rights"]["acknowledgement_required"] and not p.rights_acknowledged:
        raise ValidationFailed(
            "Some declared rights are unknown or restricted; acknowledge them to download."
        )
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for path, data in sorted(files.items()):
            info = zipfile.ZipInfo(path, date_time=EPOCH)
            info.create_system = 3
            info.external_attr = 0o100444 << 16
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, data)
    return stream.getvalue(), sha256, manifest
