"""DOC-004 style rights review and deterministic XLSX for coherent export packages.

Expected rights are the declared statuses this test registers itself.
"""

import base64
import io
import time
import zipfile

import pytest
from openpyxl import load_workbook

from specmate.export_packages import operations  # noqa: F401
from specmate.export_packages.verification import verify_archive
from specmate.operations import ValidationFailed

from .conftest import op
from .scope_helpers import workspace_of
from .test_catalog import identifier_payload, rep_payload
from .test_export_packages import download, review, rows, system  # noqa: F401
from .test_sources import ORIGINAL, files, register


def request(system_):
    return {"root_revision_id": system_["root"]["revision_id"], "quantity": 3}


def zipped(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def test_workbook_is_deterministic_and_keeps_csv_text_safe(users, system):  # noqa: F811
    payload = request(system)
    reviewed = review(users["viewer"], payload)
    first, files_a = download(users["viewer"], payload, reviewed)
    time.sleep(1.1)
    second, files_b = download(users["viewer"], payload, reviewed)
    assert first == second
    assert files_a["spreadsheet/package.xlsx"] == files_b["spreadsheet/package.xlsx"]
    assert "spreadsheet/occurrences.csv" in files_a
    raw = files_a["spreadsheet/package.xlsx"]
    with zipfile.ZipFile(io.BytesIO(raw)) as workbook:
        assert workbook.testzip() is None
        assert all(info.date_time == (1980, 1, 1, 0, 0, 0) for info in workbook.infolist())
        assert workbook.namelist() == sorted(workbook.namelist())
    book = load_workbook(io.BytesIO(raw))
    assert book.properties.created.year == 1980 and book.properties.modified.year == 1980
    assert book.sheetnames == ["occurrences", "revision-quantities"]
    sheet = book["occurrences"]
    header = [c.value for c in sheet[1]]
    table = [
        dict(zip(header, [c.value for c in line], strict=True))
        for line in sheet.iter_rows(min_row=2)
    ]
    controller = next(r for r in table if r["path"] == "controller")
    # The hostile display ref is a string cell, not a formula.
    assert controller["display_ref"] == "=HYPERLINK(1)"
    assert sheet.cell(row=1 + table.index(controller) + 1, column=2).data_type == "s"
    assert controller["scope_quantity"] == 3
    # Cells agree with the CSV for the same rows.
    csv_rows = rows(files_a["spreadsheet/occurrences.csv"])
    assert [r["path"] for r in csv_rows] == [r["path"] for r in table]
    assert not verify_archive(first, reviewed["sha256"])
    assert "spreadsheet.xlsx" not in {u["id"] for u in reviewed["manifest"]["unsupported"]}


def test_unknown_quantity_is_an_empty_workbook_cell(users, system):  # noqa: F811
    payload = {"root_revision_id": system["root"]["revision_id"]}
    reviewed = review(users["viewer"], payload)
    _, package = download(users["viewer"], payload, reviewed)
    sheet = load_workbook(io.BytesIO(package["spreadsheet/package.xlsx"]))["revision-quantities"]
    header = [c.value for c in sheet[1]]
    cable = next(
        line
        for line in sheet.iter_rows(min_row=2, values_only=True)
        if line[0] == system["cable"]["revision_id"]
    )
    row = dict(zip(header, cable, strict=True))
    assert row["scope_quantity"] is None and row["scope_quantity_status"] == "unknown"


def test_rights_review_lists_declared_status_and_gates_download(users, project, system):  # noqa: F811
    editor = users["editor"]
    board = system["board"]
    register(editor, project, ORIGINAL)
    source = op(
        editor,
        "source.create_original_bundle",
        {"workspace_id": workspace_of(project), "name": "Symbol source", "files": files(ORIGINAL)},
    )
    cleared = op(editor, "catalog.register_representation", rep_payload(board, source))
    restricted = op(
        editor,
        "catalog.register_source_identifier",
        identifier_payload(board, "Synthetic:Restricted", "Vendor terms", "restricted"),
    )
    unknown = op(
        editor,
        "catalog.register_source_identifier",
        identifier_payload(board, "Synthetic:Unknown", "unknown", "unknown"),
    )
    payload = request(system)
    reviewed = review(users["viewer"], payload)
    rights = reviewed["manifest"]["rights"]
    by_subject = {r["subject"]: r for r in rights["items"]}
    declared = {
        s: r["rights_status"] for s, r in by_subject.items() if not s.startswith("revision:")
    }
    assert declared == {
        f"representation:{cleared['representation_id']}": "cleared",
        f"source_identifier:{restricted['source_identifier_id']}": "restricted",
        f"source_identifier:{unknown['source_identifier_id']}": "unknown",
    }
    # Every revision without a declaration is an explicit unknown row with the reason.
    undeclared = rights["undeclared_revision_ids"]
    assert undeclared
    for rid in undeclared:
        row = by_subject[f"revision:{rid}"]
        assert row["rights_status"] == "unknown" and row["label"] == "no rights declared"
    assert rights["counts"] == {"cleared": 1, "restricted": 1, "unknown": 1 + len(undeclared)}
    assert rights["acknowledgement_required"] is True
    assert rights["distribution"] == "not_authorized"
    # Revisions with no declaration are listed, never treated as cleared.
    assert board["revision_id"] not in rights["undeclared_revision_ids"]
    assert system["cable"]["revision_id"] in rights["undeclared_revision_ids"]
    assert reviewed["manifest"]["sections"]["rights"]["restricted"] == 1

    body = {**payload, "reviewed_sha256": reviewed["sha256"]}
    with pytest.raises(ValidationFailed, match="acknowledge"):
        op(users["viewer"], "export_package.download", body)
    raw, package = download(users["viewer"], {**payload, "rights_acknowledged": True}, reviewed)
    assert "rights/review.json" in package and "rights/review.csv" in package
    csv_rows = rows(package["rights/review.csv"])
    assert sorted(r["rights_status"] for r in csv_rows) == (
        ["cleared", "restricted"] + ["unknown"] * (1 + len(undeclared))
    )
    assert not verify_archive(raw, reviewed["sha256"])
    # The acknowledgement does not change the reviewed manifest hash.
    again = op(users["viewer"], "export_package.download", {**body, "rights_acknowledged": True})
    assert base64.b64decode(again["archive_base64"]) == raw


def test_all_cleared_rights_need_no_acknowledgement(users, project, system):  # noqa: F811
    editor = users["editor"]
    op(
        editor,
        "catalog.register_source_identifier",
        identifier_payload(system["board"], "Synthetic:Cleared", "Open", "cleared"),
    )
    payload = request(system)
    # Undeclared revisions count as unknown and require the acknowledgement.
    first = review(users["viewer"], payload)["manifest"]["rights"]
    assert first["acknowledgement_required"] is True and first["undeclared_revision_ids"]
    for rid in first["undeclared_revision_ids"]:
        op(
            editor,
            "catalog.register_source_identifier",
            identifier_payload({"revision_id": rid}, f"Synthetic:{rid}", "Open", "cleared"),
        )
    reviewed = review(users["viewer"], payload)
    assert reviewed["manifest"]["rights"]["acknowledgement_required"] is False
    body = {**payload, "reviewed_sha256": reviewed["sha256"]}
    result = op(users["viewer"], "export_package.download", body)
    assert result["archive_base64"]
